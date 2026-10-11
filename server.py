"""
Engineering Progress Tracking Dashboard Server
Reads the master Excel MDR file and serves a web dashboard with:
- Weekly progress tracking
- Document status call-outs
- 2-week look-ahead warnings
- S-Curve visualization data
Supports file upload to replace the master Excel file.
"""

import http.server
import socketserver
import json
import os
import re
import datetime
import threading
import shutil
import tempfile
import io
import warnings
import gc
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

# Suppress harmless openpyxl template/formatting warnings
warnings.filterwarnings('ignore', category=UserWarning, module='openpyxl')

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
except ImportError:
    print("Installing openpyxl...")
    import subprocess
    subprocess.check_call(["pip", "install", "openpyxl"])
    import openpyxl
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ─── Configuration ───────────────────────────────────────────────────────────

BASE_DIR = Path(__file__).parent
PORT = int(os.environ.get("PORT", 8081))
CACHE_FILE = BASE_DIR / "data_cache.json"

# Work Package sheet configurations
WP_SHEETS = [
    {"sheet": "WP01 TOPSIDE", "label": "WP01 Topside"},
    {"sheet": "WP01 JACKET",  "label": "WP01 Jacket"},
    {"sheet": "WP02 TOPSIDE", "label": "WP02 Topside"},
    {"sheet": "WP02 JACKET",  "label": "WP02 Jacket"},
]

SCURVE_SHEETS = [
    {"sheet": "Eng. S-Curve_OVERALL",     "label": "Overall"},
    {"sheet": "Eng. S-Curve_WP01 TOPSIDE","label": "WP01 Topside"},
    {"sheet": "Eng. S-Curve_WP01 JACKET", "label": "WP01 Jacket"},
    {"sheet": "Eng. S-Curve_WP02 TOPSIDE","label": "WP02 Topside"},
    {"sheet": "Eng. S-Curve_WP02 JACKET", "label": "WP02 Jacket"},
]

NEEDED_SHEET_NAMES = {
    "Sum",
    "Weekly Summary",
    "WP01 TOPSIDE",
    "WP01 JACKET",
    "WP02 TOPSIDE",
    "WP02 JACKET",
    "PRO.ENGINEERING_MR TBE",
    "Overdue Summary",
    "Overdue List",
    "Eng. S-Curve_OVERALL",
    "Eng. S-Curve_WP01 TOPSIDE",
    "Eng. S-Curve_WP01 JACKET",
    "Eng. S-Curve_WP02 TOPSIDE",
    "Eng. S-Curve_WP02 JACKET",
}

# Data cache
_data_cache = {
    "data": None, 
    "lock": threading.Lock(), 
    "file": None,
    "is_extracting": False
}


# ─── File Detection ──────────────────────────────────────────────────────────

MONTH_MAP = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12
}

def parse_date_from_filename(filename):
    """Extract date from filename like 'Cut off 03-Jul-26' or similar patterns."""
    # Pattern: DD-Mon-YY (e.g., 03-Jul-26)
    match = re.search(r'(\d{1,2})[-\s](Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[-\s](\d{2,4})', filename, re.IGNORECASE)
    if match:
        day = int(match.group(1))
        month = MONTH_MAP.get(match.group(2).lower(), 1)
        year = int(match.group(3))
        if year < 100:
            year += 2000
        try:
            return datetime.date(year, month, day)
        except ValueError:
            pass
    # Pattern: YYYYMMDD
    match = re.search(r'(\d{4})(\d{2})(\d{2})', filename)
    if match:
        try:
            return datetime.date(int(match.group(1)), int(match.group(2)), int(match.group(3)))
        except ValueError:
            pass
    return None


def find_latest_excel():
    """Find the latest Excel file in BASE_DIR by parsing dates from filenames."""
    excel_files = list(BASE_DIR.glob("*.xlsx"))
    if not excel_files:
        return None

    # Try to find files with dates and pick the latest
    dated_files = []
    for f in excel_files:
        if f.name.startswith("~$"):
            continue  # Skip temp files
        if "MDR" not in f.name.upper():
            continue  # Skip non-MDR files (e.g. Procurement Plans)
            
        dt = parse_date_from_filename(f.name)
        if dt:
            dated_files.append((dt, f))

    if dated_files:
        dated_files.sort(key=lambda x: x[0], reverse=True)
        return dated_files[0][1]

    # Fallback: return the most recently modified xlsx
    valid = [f for f in excel_files if not f.name.startswith("~$")]
    if valid:
        valid.sort(key=lambda f: f.stat().st_mtime, reverse=True)
        return valid[0]
    return None


# ─── Caching & Memory Optimization ──────────────────────────────────────────

def get_file_fingerprint(filepath):
    """Return a unique fingerprint based on filename, size, and mtime."""
    p = Path(filepath)
    if not p.exists():
        return None
    stat = p.stat()
    return f"{p.name}_{stat.st_size}_{int(stat.st_mtime)}"


def load_disk_cache(excel_file):
    """Load cached JSON data if it exists and matches the Excel file fingerprint."""
    if not CACHE_FILE.exists() or not excel_file:
        return None
    try:
        current_fp = get_file_fingerprint(excel_file)
        with open(CACHE_FILE, "r", encoding="utf-8") as f:
            cached_payload = json.load(f)
        if cached_payload.get("_fingerprint") == current_fp:
            print(f"  [Cache] Loaded precomputed JSON cache from {CACHE_FILE.name} (Instant cold-boot)")
            return cached_payload.get("data")
    except Exception as e:
        print(f"  [Cache] Warning: Could not read cache: {e}")
    return None


def save_disk_cache(excel_file, data):
    """Save extracted data to disk cache."""
    if not excel_file or not data:
        return
    try:
        current_fp = get_file_fingerprint(excel_file)
        payload = {
            "_fingerprint": current_fp,
            "cached_at": datetime.datetime.now().isoformat(),
            "data": data
        }
        with open(CACHE_FILE, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, default=str)
        print(f"  [Cache] Saved extracted JSON cache to {CACHE_FILE.name}")
    except Exception as e:
        print(f"  [Cache] Warning: Could not write cache: {e}")


def load_trimmed_workbook(excel_path):
    """
    Load an Excel workbook keeping ONLY the sheets required by the dashboard.
    Bypasses huge unused raw database sheets (such as MDR with 22MB XML and calcChain with 17MB XML),
    reducing peak memory from >430MB down to ~230MB to stay safely within 512MB RAM environments (e.g. Render).
    """
    excel_path = Path(excel_path)
    t0 = datetime.datetime.now()

    with zipfile.ZipFile(excel_path, 'r') as src_zip:
        file_list = src_zip.namelist()

        # Check if workbook.xml exists
        if 'xl/workbook.xml' not in file_list or 'xl/_rels/workbook.xml.rels' not in file_list:
            return openpyxl.load_workbook(str(excel_path), data_only=True)

        wb_xml = src_zip.read('xl/workbook.xml')
        tree = ET.fromstring(wb_xml)
        ns = {
            'main': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
            'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
        }

        rels_xml = src_zip.read('xl/_rels/workbook.xml.rels')
        rels_tree = ET.fromstring(rels_xml)
        rel_ns = {'rel': 'http://schemas.openxmlformats.org/package/2006/relationships'}

        rid_to_target = {}
        for rel in rels_tree.findall('.//rel:Relationship', rel_ns):
            rid_to_target[rel.get('Id')] = rel.get('Target')

        keep_files = set()
        sheets_node = tree.find('.//main:sheets', ns)
        if sheets_node is not None:
            for s in list(sheets_node):
                name = s.get('name')
                rid = s.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
                target = rid_to_target.get(rid, '')
                if name in NEEDED_SHEET_NAMES:
                    keep_files.add('xl/' + target)
                else:
                    sheets_node.remove(s)

        new_wb_xml = ET.tostring(tree, xml_declaration=True, encoding='utf-8')

        out_bio = io.BytesIO()
        with zipfile.ZipFile(out_bio, 'w', compression=zipfile.ZIP_DEFLATED) as dst_zip:
            for item in src_zip.infolist():
                if item.filename == 'xl/workbook.xml':
                    dst_zip.writestr(item, new_wb_xml)
                elif item.filename.startswith('xl/worksheets/'):
                    if item.filename in keep_files:
                        dst_zip.writestr(item, src_zip.read(item.filename))
                elif item.filename == 'xl/calcChain.xml':
                    # Omit massive formula dependency chain (saves 17MB XML)
                    continue
                else:
                    dst_zip.writestr(item, src_zip.read(item.filename))

        out_bio.seek(0)
        wb = openpyxl.load_workbook(out_bio, data_only=True)
        dur = (datetime.datetime.now() - t0).total_seconds()
        print(f"  [Memory Opt] Loaded trimmed workbook (14 sheets) in {dur:.2f}s")
        return wb


# ─── Helper Functions ─────────────────────────────────────────────────────────

def safe_str(val):
    """Convert value to string safely."""
    if val is None:
        return ""
    return str(val).strip()


def safe_float(val):
    """Convert value to float safely."""
    if val is None:
        return None
    if isinstance(val, str):
        if val in ("#REF!", "#N/A", "#VALUE!", "N/A", "-", ""):
            return None
        try:
            return float(val.replace(",", "").replace("%", ""))
        except ValueError:
            return None
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def safe_date(val):
    """Convert value to ISO date string safely."""
    if val is None:
        return None
    if isinstance(val, datetime.datetime):
        # Filter out Excel zero-dates (1900-01-00 or time(0,0))
        if val.year < 1950:
            return None
        return val.strftime("%Y-%m-%d")
    if isinstance(val, datetime.date):
        if val.year < 1950:
            return None
        return val.isoformat()
    if isinstance(val, datetime.time):
        # time(0,0) means empty in this Excel
        return None
    s = str(val).strip()
    if not s or s in ("0", "N/A", "-", "None", "#N/A", "#REF!", "#VALUE!"):
        return None
    # Validate ISO date string (YYYY-MM-DD or YYYY-MM-DD HH:MM:SS)
    iso_match = re.match(r'^(\d{4})-(\d{2})-(\d{2})', s)
    if iso_match:
        try:
            y, m, d = int(iso_match.group(1)), int(iso_match.group(2)), int(iso_match.group(3))
            if 1950 <= y <= 2100:
                return f"{y:04d}-{m:02d}-{d:02d}"
        except ValueError:
            pass
    # Pattern: DD-Mon-YY or DD-Mon-YYYY
    text_match = re.search(r'^(\d{1,2})[-\s](Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[-\s](\d{2,4})$', s, re.IGNORECASE)
    if text_match:
        try:
            day = int(text_match.group(1))
            month = MONTH_MAP.get(text_match.group(2).lower(), 1)
            year = int(text_match.group(3))
            if year < 100:
                year += 2000
            if 1950 <= year <= 2100:
                dt = datetime.date(year, month, day)
                return dt.isoformat()
        except ValueError:
            pass
    return None


def safe_code(val):
    """Convert return code value to clean text string safely."""
    if val is None:
        return None
    if isinstance(val, (datetime.time, datetime.datetime, datetime.date)):
        return None
    s = str(val).strip()
    if not s or s in ("-", "None", "00:00:00", "0", "N/A", "#N/A"):
        return None
    return s


def cell_val(ws, row, col):
    """Get cell value by row number and column letter."""
    return ws[f"{col}{row}"].value


def col_letter_to_idx(col_str):
    """Convert column letter(s) to 0-based index."""
    result = 0
    for char in col_str.upper():
        result = result * 26 + (ord(char) - ord('A') + 1)
    return result - 1


# ─── Data Extraction ──────────────────────────────────────────────────────────

def extract_summary(wb, filename=None):
    """Extract overall summary from 'Sum' sheet."""
    ws = wb["Sum"]
    cutoff_date = safe_date(ws["D9"].value)
    prev_cutoff = safe_date(ws["K9"].value)

    # Dynamic row detection based on Column C labels (to handle row insertions like Rebaseline Variance)
    row_map = {
        "plan": 12, "forecast": 13, "rebaseline_variance": None, "actual": 14,
        "variance": 15, "spi_weekly": 16, "spi_cum": 17, "area_concern": 18
    }
    for r in range(12, 25):
        label = safe_str(ws.cell(row=r, column=3).value)
        if not label: continue
        label_up = label.upper().strip()
        if label_up == "LEGEND:": break # Stop at legend to avoid matching legend text
        if label_up == "PLAN PROGRESS": row_map["plan"] = r
        elif label_up in ["FORECAST PROGRESS", "REBASELINE PROGRESS"]: row_map["forecast"] = r
        elif label_up == "REBASELINE VARIANCE": row_map["rebaseline_variance"] = r
        elif label_up == "ACTUAL PROGRESS": row_map["actual"] = r
        elif label_up == "VARIANCE": row_map["variance"] = r
        elif label_up == "SPI (WEEKLY)": row_map["spi_weekly"] = r
        elif label_up == "SPI (CUMULATIVE)": row_map["spi_cum"] = r
        elif label_up == "AREA OF CONCERN": row_map["area_concern"] = r

    r_plan, r_fc, r_act = row_map["plan"], row_map["forecast"], row_map["actual"]
    r_reb_var = row_map.get("rebaseline_variance")
    r_var, r_spiw, r_spic = row_map["variance"], row_map["spi_weekly"], row_map["spi_cum"]
    r_aoc = row_map["area_concern"]

    plan_last = safe_float(ws[f"D{r_plan}"].value) if r_plan else None
    plan_this = safe_float(ws[f"E{r_plan}"].value) if r_plan else None
    fc_last = safe_float(ws[f"D{r_fc}"].value) if r_fc else None
    fc_this = safe_float(ws[f"E{r_fc}"].value) if r_fc else None
    act_last = safe_float(ws[f"D{r_act}"].value) if r_act else None
    act_this = safe_float(ws[f"E{r_act}"].value) if r_act else None

    # Rebased Line Variance (+0.30% in B1: Actual - Rebaseline/Forecast)
    reb_var_last = safe_float(ws[f"D{r_reb_var}"].value) if r_reb_var else None
    reb_var_this = safe_float(ws[f"E{r_reb_var}"].value) if r_reb_var else None
    if reb_var_this is None and act_this is not None and fc_this is not None:
        reb_var_this = round(act_this - fc_this, 6)
    if reb_var_last is None and act_last is not None and fc_last is not None:
        reb_var_last = round(act_last - fc_last, 6)

    # Non-Rebased Line Variance (-5.46% in B1: Actual - Original Baseline Plan)
    non_reb_var_this = round(act_this - plan_this, 6) if (act_this is not None and plan_this is not None) else None
    non_reb_var_last = round(act_last - plan_last, 6) if (act_last is not None and plan_last is not None) else None

    return {
        "cutoff_date": cutoff_date,
        "prev_cutoff_date": prev_cutoff,
        "plan_last_week": plan_last,
        "plan_this_week": plan_this,
        "forecast_last_week": fc_last,
        "forecast_this_week": fc_this,
        "actual_last_week": act_last,
        "actual_this_week": act_this,
        "variance_rebased_last": reb_var_last,
        "variance_rebased_this": reb_var_this,
        "variance_non_rebased_last": non_reb_var_last,
        "variance_non_rebased_this": non_reb_var_this,
        "variance_last_week": reb_var_last if reb_var_last is not None else non_reb_var_last,
        "variance_this_week": reb_var_this if reb_var_this is not None else non_reb_var_this,
        "spi_weekly_last": safe_float(ws[f"D{r_spiw}"].value) if r_spiw else None,
        "spi_weekly_this": safe_float(ws[f"E{r_spiw}"].value) if r_spiw else None,
        "spi_cumulative_last": safe_float(ws[f"D{r_spic}"].value) if r_spic else None,
        "spi_cumulative_this": safe_float(ws[f"E{r_spic}"].value) if r_spic else None,
        "area_of_concern": safe_str(ws[f"G{r_aoc}"].value) if r_aoc else "",
        "project_name": "EPC OF WELLHEAD PLATFORMS FOR AUNG SINKHA DEVELOPMENT PROJECT PHASE 1A (EPC-01)",
        "contractor": "GC Maintenance and Engineering Company Limited",
        "job_no": "SD-20-26600-01",
    }


def extract_weekly_summary(wb, filename=None):
    """Extract discipline-level weekly progress from 'Weekly Summary' sheet."""
    ws = wb["Weekly Summary"]
    cutoff_date = safe_date(ws["U7"].value)

    disciplines = []
    current_wp = None

    for row in range(12, 45):
        c_val = safe_str(ws[f"C{row}"].value)
        d_val = safe_str(ws[f"D{row}"].value)

        # Work package header row
        if c_val and c_val.startswith("WP"):
            current_wp = c_val
            wp_data = {
                "name": c_val,
                "is_wp_header": True,
                "weight_l1": safe_float(ws[f"F{row}"].value),
                "weight_l2": safe_float(ws[f"G{row}"].value),
                "last_plan": safe_float(ws[f"H{row}"].value),
                "last_forecast": safe_float(ws[f"I{row}"].value),
                "last_actual_medms": safe_float(ws[f"J{row}"].value),
                "last_actual": safe_float(ws[f"K{row}"].value),
                "incr_plan": safe_float(ws[f"L{row}"].value),
                "incr_forecast": safe_float(ws[f"M{row}"].value),
                "incr_actual_medms": safe_float(ws[f"N{row}"].value),
                "incr_actual": safe_float(ws[f"O{row}"].value),
                "this_plan": safe_float(ws[f"P{row}"].value),
                "this_forecast": safe_float(ws[f"Q{row}"].value),
                "this_actual_medms": safe_float(ws[f"R{row}"].value),
                "this_actual": safe_float(ws[f"S{row}"].value),
                "var_actual_plan": safe_float(ws[f"T{row}"].value),
                "var_actual_forecast": safe_float(ws[f"U{row}"].value),
            }
            disciplines.append(wp_data)
            continue

        # Overall engineering progress row
        if c_val and "OVERALL" in c_val.upper():
            overall = {
                "name": c_val,
                "is_wp_header": True,
                "is_overall": True,
                "weight_l1": safe_float(ws[f"F{row}"].value),
                "last_plan": safe_float(ws[f"H{row}"].value),
                "last_actual": safe_float(ws[f"K{row}"].value),
                "incr_plan": safe_float(ws[f"L{row}"].value),
                "incr_actual": safe_float(ws[f"O{row}"].value),
                "this_plan": safe_float(ws[f"P{row}"].value),
                "this_forecast": safe_float(ws[f"Q{row}"].value),
                "this_actual": safe_float(ws[f"S{row}"].value),
                "var_actual_plan": safe_float(ws[f"T{row}"].value),
                "var_actual_forecast": safe_float(ws[f"U{row}"].value),
            }
            disciplines.append(overall)
            continue

        # Discipline detail row
        if d_val and current_wp:
            disc = {
                "name": d_val,
                "wp": current_wp,
                "is_wp_header": False,
                "weight_l2": safe_float(ws[f"G{row}"].value),
                "last_plan": safe_float(ws[f"H{row}"].value),
                "last_forecast": safe_float(ws[f"I{row}"].value),
                "last_actual": safe_float(ws[f"K{row}"].value),
                "incr_plan": safe_float(ws[f"L{row}"].value),
                "incr_forecast": safe_float(ws[f"M{row}"].value),
                "incr_actual": safe_float(ws[f"O{row}"].value),
                "this_plan": safe_float(ws[f"P{row}"].value),
                "this_forecast": safe_float(ws[f"Q{row}"].value),
                "this_actual": safe_float(ws[f"S{row}"].value),
                "var_actual_plan": safe_float(ws[f"T{row}"].value),
                "var_actual_forecast": safe_float(ws[f"U{row}"].value),
            }
            disciplines.append(disc)

    # SPI
    spi_weekly = safe_float(ws["O44"].value)
    spi_cumulative = safe_float(ws["T44"].value)

    return {
        "cutoff_date": cutoff_date,
        "disciplines": disciplines,
        "spi_weekly": spi_weekly,
        "spi_cumulative": spi_cumulative,
    }


def detect_milestone_columns(ws, header_group_row, subheader_row, start_col=20, max_col=65):
    """
    Dynamically scan milestone group headers (e.g. IFR, IFA, AFC, AP or IFA (A1), IFA (C1))
    and their subheaders (PLAN, REBASELINE PLAN, FORECAST, D2 SUBMIT ID, SUBMIT TO PTTEPI DATE, etc.)
    to map column numbers for each milestone block, avoiding errors caused by newly inserted template columns.
    """
    group_spans = []
    for c in range(start_col, max_col + 1):
        val = ws.cell(row=header_group_row, column=c).value
        if val:
            v_str = str(val).strip().upper()
            # Match recognized milestone group labels
            for name in ["IFA (A1)", "IFA (C1)", "IFR", "IFA", "AFC", "AP"]:
                if v_str == name or v_str.startswith(name):
                    group_spans.append((c, name))
                    break

    milestones = {}
    for idx, (c_start, ms_name) in enumerate(group_spans):
        c_end = group_spans[idx + 1][0] - 1 if idx + 1 < len(group_spans) else max_col
        key = ms_name
        if key == "AP":
            key = f"AP_{idx}"
        milestones[key] = {}
        for c in range(c_start, c_end + 1):
            sub = ws.cell(row=subheader_row, column=c).value
            if not sub:
                continue
            sub_norm = re.sub(r'\s+', ' ', str(sub).strip().upper())
            if "SUBMIT TO PTTEPI" in sub_norm:
                milestones[key]["submit_date"] = c
            elif "SUBMIT ID" in sub_norm or "D2 SUBMIT" in sub_norm:
                milestones[key]["submit_id"] = c
            elif "REBASELINE" in sub_norm:
                milestones[key]["rebaseline"] = c
            elif "FORECAST" in sub_norm:
                milestones[key]["forecast"] = c
            elif "PLAN" in sub_norm and "REBASELINE" not in sub_norm:
                milestones[key]["plan"] = c
            elif "RETURN FROM PTTEPI" in sub_norm or "RETURN DATE" in sub_norm:
                milestones[key]["return_date"] = c
            elif "RETURN CODE" in sub_norm:
                milestones[key]["return_code"] = c
            elif "ACTUAL" in sub_norm:
                milestones[key]["actual"] = c
            elif "TRANSMITTAL" in sub_norm:
                milestones[key]["transmittal"] = c

    return milestones


def extract_documents(wb, sheet_name, wp_label):
    """Extract document-level data from a WP sheet."""
    if sheet_name not in wb.sheetnames:
        return []

    ws = wb[sheet_name]
    documents = []

    # Dynamically detect milestone columns from Rows 12 & 13
    col_map = detect_milestone_columns(ws, header_group_row=12, subheader_row=13)
    ifr_cols = col_map.get("IFR", {})
    ifa_cols = col_map.get("IFA", {})
    afc_cols = col_map.get("AFC", {})
    ap_cols = next((v for k, v in col_map.items() if k.startswith("AP")), {})

    def get_c(r, dyn_dict, field_key, fallback_letter):
        c_idx = dyn_dict.get(field_key)
        if c_idx:
            return ws.cell(row=r, column=c_idx).value
        if fallback_letter:
            return ws[f"{fallback_letter}{r}"].value
        return None

    for row in range(17, ws.max_row + 1):
        doc_no = safe_str(ws[f"I{row}"].value)
        if not doc_no or doc_no.startswith("#"):
            continue

        b_val = ws[f"B{row}"].value
        if b_val is None:
            continue
        # Skip sub-header / section rows
        try:
            int(b_val)
        except (ValueError, TypeError):
            continue

        # IFR dates
        ifr_plan = safe_date(get_c(row, ifr_cols, "plan", "V"))
        ifr_forecast = safe_date(get_c(row, ifr_cols, "forecast", None) or get_c(row, ifr_cols, "rebaseline", "W"))
        ifr_submit_id = safe_str(get_c(row, ifr_cols, "submit_id", "X"))
        ifr_submit_date = safe_date(get_c(row, ifr_cols, "submit_date", "Y"))
        ifr_actual = safe_date(get_c(row, ifr_cols, "actual", "Z"))
        ifr_transmittal = safe_str(get_c(row, ifr_cols, "transmittal", "AA"))

        # IFA dates
        ifa_plan = safe_date(get_c(row, ifa_cols, "plan", "AB"))
        ifa_forecast = safe_date(get_c(row, ifa_cols, "forecast", None) or get_c(row, ifa_cols, "rebaseline", "AC"))
        ifa_submit_id = safe_str(get_c(row, ifa_cols, "submit_id", "AD"))
        ifa_submit_date = safe_date(get_c(row, ifa_cols, "submit_date", "AE"))
        ifa_actual = safe_date(get_c(row, ifa_cols, "actual", "AF"))
        ifa_transmittal = safe_str(get_c(row, ifa_cols, "transmittal", "AG"))

        # AFC dates
        afc_plan = safe_date(get_c(row, afc_cols, "plan", "AH"))
        afc_forecast = safe_date(get_c(row, afc_cols, "forecast", None) or get_c(row, afc_cols, "rebaseline", "AI"))
        afc_submit_id = safe_str(get_c(row, afc_cols, "submit_id", "AJ"))
        afc_submit_date = safe_date(get_c(row, afc_cols, "submit_date", "AK"))
        afc_actual = safe_date(get_c(row, afc_cols, "actual", "AL"))
        afc_transmittal = safe_str(get_c(row, afc_cols, "transmittal", "AM"))

        # AP (Approval/Return)
        ap_plan = safe_date(get_c(row, ap_cols, "plan", "AN"))
        ap_forecast = safe_date(get_c(row, ap_cols, "forecast", None) or get_c(row, ap_cols, "rebaseline", "AO"))
        ap_return_code = safe_code(get_c(row, ap_cols, "return_code", "AP"))
        ap_return_date = safe_date(get_c(row, ap_cols, "return_date", "AQ"))

        # Determine status based on "Submit to PTTEPI Date" columns
        if afc_submit_date and afc_submit_date not in ("", None):
            status = "AFC Submitted"
        elif ifa_submit_date and ifa_submit_date not in ("", None):
            status = "IFA Submitted"
        elif ifr_submit_date and ifr_submit_date not in ("", None):
            status = "IFR Submitted"
        else:
            status = "Not Yet Submitted"

        # Revision scheme / Remarks
        remarks = safe_str(ws[f"U{row}"].value)

        # PENDING FINAL APPROVAL condition for Regular Eng Deliverables: AFC submit date filled AND AP Return date blank
        is_pending_final_approval = bool(afc_submit_date) and not bool(ap_return_date)

        doc = {
            "no": b_val,
            "area": safe_str(ws[f"C{row}"].value),
            "discipline": safe_str(ws[f"D{row}"].value),
            "type": safe_str(ws[f"E{row}"].value),
            "seq_no": safe_str(ws[f"F{row}"].value),
            "format": safe_str(ws[f"G{row}"].value),
            "doc_no": doc_no,
            "title": safe_str(ws[f"J{row}"].value),
            "class": safe_str(ws[f"K{row}"].value),
            "mh": safe_float(ws[f"L{row}"].value),
            "weight_l2": safe_float(ws[f"N{row}"].value),
            "plan_pct": safe_float(ws[f"P{row}"].value),
            "forecast_pct": safe_float(ws[f"Q{row}"].value),
            "actual_pct": safe_float(ws[f"R{row}"].value),
            "variance": safe_float(ws[f"T{row}"].value),
            "revision_scheme": remarks,
            "wp": wp_label,
            "status": status,
            "is_mr_tbe": False,
            "is_pending_final_approval": is_pending_final_approval,
            "is_cy_ap_pending": is_pending_final_approval,
            "is_complete": bool(ap_return_date),
            # IFR
            "ifr_plan": ifr_plan,
            "ifr_forecast": ifr_forecast,
            "ifr_submit_id": ifr_submit_id,
            "ifr_submit_date": ifr_submit_date,
            "ifr_actual": ifr_actual,
            "ifr_transmittal": ifr_transmittal,
            # IFA
            "ifa_plan": ifa_plan,
            "ifa_forecast": ifa_forecast,
            "ifa_submit_id": ifa_submit_id,
            "ifa_submit_date": ifa_submit_date,
            "ifa_actual": ifa_actual,
            "ifa_transmittal": ifa_transmittal,
            # AFC
            "afc_plan": afc_plan,
            "afc_forecast": afc_forecast,
            "afc_submit_id": afc_submit_id,
            "afc_submit_date": afc_submit_date,
            "afc_actual": afc_actual,
            "afc_transmittal": afc_transmittal,
            # AP
            "ap_plan": ap_plan,
            "ap_forecast": ap_forecast,
            "ap_return_code": ap_return_code,
            "ap_return_date": ap_return_date,
        }
        documents.append(doc)

    return documents


def extract_mr_tbe_documents(wb):
    """Extract procurement engineering MR/TBE documents from 'PRO.ENGINEERING_MR TBE' sheet."""
    if "PRO.ENGINEERING_MR TBE" not in wb.sheetnames:
        return []

    ws = wb["PRO.ENGINEERING_MR TBE"]
    documents = []
    current_wp = "WP01 Topside"

    col_map = detect_milestone_columns(ws, header_group_row=10, subheader_row=11)
    a1_cols = col_map.get("IFA (A1)", {})
    c1_cols = col_map.get("IFA (C1)", {})
    ap_groups = [v for k, v in col_map.items() if k.startswith("AP")]
    ap1_cols = ap_groups[0] if len(ap_groups) > 0 else {}
    ap2_cols = ap_groups[1] if len(ap_groups) > 1 else {}

    def get_mr_c(r, dyn_dict, field_key, fallback_col_idx):
        c_idx = dyn_dict.get(field_key)
        if c_idx:
            return ws.cell(row=r, column=c_idx).value
        if fallback_col_idx:
            return ws.cell(row=r, column=fallback_col_idx).value
        return None

    for r in range(12, ws.max_row + 1):
        b_val = ws.cell(row=r, column=2).value  # Column B
        if not b_val:
            continue
        
        b_str = str(b_val).strip()
        if b_str.startswith("WP01 TOPSIDE"):
            current_wp = "WP01 Topside"
            continue
        elif b_str.startswith("WP01 JACKET"):
            current_wp = "WP01 Jacket"
            continue
        elif b_str.startswith("WP02 TOPSIDE"):
            current_wp = "WP02 Topside"
            continue
        elif b_str.startswith("WP02 JACKET"):
            current_wp = "WP02 Jacket"
            continue

        try:
            int(b_val)
        except (ValueError, TypeError):
            continue

        doc_no = safe_str(ws.cell(row=r, column=9).value)  # Column I
        if not doc_no or doc_no.startswith("#"):
            continue

        disc = safe_str(ws.cell(row=r, column=4).value)      # Column D: Disc
        doc_type = safe_str(ws.cell(row=r, column=5).value)  # Column E: Type (MR / TBE)
        title = safe_str(ws.cell(row=r, column=10).value)    # Column J: Document Title
        remarks = safe_str(ws.cell(row=r, column=21).value)  # Column U: Remarks / Revision Scheme

        # Milestone 1: IFA (A1)
        ifa_plan = safe_date(get_mr_c(r, a1_cols, "plan", 22))
        ifa_forecast = safe_date(get_mr_c(r, a1_cols, "forecast", None) or get_mr_c(r, a1_cols, "rebaseline", 23))
        ifa_submit_id = safe_str(get_mr_c(r, a1_cols, "submit_id", 24))
        ifa_submit_date = safe_date(get_mr_c(r, a1_cols, "submit_date", 25))
        ifa_actual = safe_date(get_mr_c(r, a1_cols, "actual", 26))
        ifa_transmittal = safe_str(get_mr_c(r, a1_cols, "transmittal", 27))

        # Milestone 1 Return: AP
        ap_plan = safe_date(get_mr_c(r, ap1_cols, "plan", 28))
        ap_forecast = safe_date(get_mr_c(r, ap1_cols, "forecast", None) or get_mr_c(r, ap1_cols, "rebaseline", 29))
        ap_return_code = safe_code(get_mr_c(r, ap1_cols, "return_code", 30))
        ap_return_date = safe_date(get_mr_c(r, ap1_cols, "return_date", 31))

        # Milestone 2: IFA (C1)
        afc_plan = safe_date(get_mr_c(r, c1_cols, "plan", 34))
        afc_forecast = safe_date(get_mr_c(r, c1_cols, "forecast", None) or get_mr_c(r, c1_cols, "rebaseline", 35))
        afc_submit_id = safe_str(get_mr_c(r, c1_cols, "submit_id", 36))
        afc_submit_date = safe_date(get_mr_c(r, c1_cols, "submit_date", 37))
        afc_actual = safe_date(get_mr_c(r, c1_cols, "actual", 38))
        afc_transmittal = safe_str(get_mr_c(r, c1_cols, "transmittal", 39))

        # Milestone 2 Return: AP
        ap2_plan = safe_date(get_mr_c(r, ap2_cols, "plan", 40))
        ap2_forecast = safe_date(get_mr_c(r, ap2_cols, "forecast", None) or get_mr_c(r, ap2_cols, "rebaseline", 41))
        ap2_return_code = safe_code(get_mr_c(r, ap2_cols, "return_code", 42))
        ap2_return_date = safe_date(get_mr_c(r, ap2_cols, "return_date", 43))

        # Determine document status
        if afc_submit_date:
            status = "IFA (C1) Submitted"
        elif ifa_submit_date:
            status = "IFA (A1) Submitted"
        else:
            status = "Not Yet Submitted"

        # PENDING FINAL APPROVAL condition: IFA (A1) submit date filled AND AP Return blank OR IFA (C1) submit date filled AND AP Return blank
        a1_pending = bool(ifa_submit_date) and not bool(ap_return_date)
        c1_pending = bool(afc_submit_date) and not bool(ap2_return_date)
        is_pending_final_approval = a1_pending or c1_pending

        doc = {
            "no": b_val,
            "area": safe_str(ws.cell(row=r, column=3).value),
            "discipline": disc,
            "type": doc_type,
            "seq_no": safe_str(ws.cell(row=r, column=6).value),
            "format": safe_str(ws.cell(row=r, column=7).value),
            "doc_no": doc_no,
            "title": title,
            "class": safe_str(ws.cell(row=r, column=11).value),
            "mh": safe_float(ws.cell(row=r, column=12).value),
            "weight_l2": safe_float(ws.cell(row=r, column=15).value),
            "plan_pct": safe_float(ws.cell(row=r, column=16).value),
            "forecast_pct": safe_float(ws.cell(row=r, column=17).value),
            "actual_pct": safe_float(ws.cell(row=r, column=18).value),
            "variance": safe_float(ws.cell(row=r, column=20).value),
            "revision_scheme": remarks,
            "wp": current_wp,
            "status": status,
            "is_mr_tbe": True,
            "is_pending_final_approval": is_pending_final_approval,
            "is_cy_ap_pending": is_pending_final_approval,
            "is_complete": bool(ap_return_date or ap2_return_date),
            # IFR (empty for MR/TBE)
            "ifr_plan": None,
            "ifr_forecast": None,
            "ifr_submit_id": None,
            "ifr_submit_date": None,
            "ifr_actual": None,
            "ifr_transmittal": None,
            # IFA (Milestone 1: IFA A1)
            "ifa_plan": ifa_plan,
            "ifa_forecast": ifa_forecast,
            "ifa_submit_id": ifa_submit_id,
            "ifa_submit_date": ifa_submit_date,
            "ifa_actual": ifa_actual,
            "ifa_transmittal": ifa_transmittal,
            # AFC (Milestone 2: IFA C1)
            "afc_plan": afc_plan,
            "afc_forecast": afc_forecast,
            "afc_submit_id": afc_submit_id,
            "afc_submit_date": afc_submit_date,
            "afc_actual": afc_actual,
            "afc_transmittal": afc_transmittal,
            # AP Returns
            "ap_plan": ap_plan,
            "ap_forecast": ap_forecast,
            "ap_return_code": ap_return_code,
            "ap_return_date": ap_return_date,
            "ap2_plan": ap2_plan,
            "ap2_forecast": ap2_forecast,
            "ap2_return_code": ap2_return_code,
            "ap2_return_date": ap2_return_date,
        }
        documents.append(doc)

    return documents


def extract_overdue_summary(wb):
    """Extract overdue summary from 'Overdue Summary' sheet."""
    if "Overdue Summary" not in wb.sheetnames:
        return {"disciplines": [], "total_text": ""}

    ws = wb["Overdue Summary"]
    total_text = safe_str(ws["B3"].value)
    progress = safe_float(ws["S3"].value)

    disciplines = []
    for row in range(6, 20):
        disc_name = safe_str(ws[f"B{row}"].value)
        if not disc_name:
            continue
        disciplines.append({
            "discipline": disc_name,
            "total": safe_float(ws[f"C{row}"].value) or 0,
            "first_rev": safe_float(ws[f"D{row}"].value) or 0,
            "ifa_afc_issued": safe_float(ws[f"E{row}"].value) or 0,
            "pending_return": safe_float(ws[f"F{row}"].value) or 0,
        })

    return {
        "total_text": total_text,
        "progress": progress,
        "disciplines": disciplines,
    }


def extract_overdue_list(wb, filename=None):
    """Extract overdue document list from 'Overdue List' sheet."""
    if "Overdue List" not in wb.sheetnames:
        return []

    ws = wb["Overdue List"]
    cutoff_date = safe_date(ws["J6"].value)
    docs = []

    for row in range(9, ws.max_row + 1):
        doc_no = safe_str(ws[f"D{row}"].value)
        if not doc_no:
            continue

        b_val = ws[f"B{row}"].value
        try:
            int(b_val)
        except (ValueError, TypeError):
            continue

        docs.append({
            "no": b_val,
            "discipline": safe_str(ws[f"C{row}"].value),
            "doc_no": doc_no,
            "rev": safe_str(ws[f"E{row}"].value),
            "description": safe_str(ws[f"F{row}"].value),
            "procurement_impact": safe_str(ws[f"G{row}"].value),
            "plan_date": safe_date(ws[f"H{row}"].value),
            "forecast_date": safe_date(ws[f"I{row}"].value),
            "remark": safe_str(ws[f"J{row}"].value),
            "ifr_submit_id": safe_str(ws[f"L{row}"].value),
            "ifr_submit_date": safe_date(ws[f"M{row}"].value),
            "ifr_due_return": safe_date(ws[f"N{row}"].value),
            "ifr_return_code": safe_str(ws[f"O{row}"].value),
            "ifr_return_date": safe_date(ws[f"P{row}"].value),
        })

    return docs


def extract_scurve(wb, sheet_name, label):
    """Extract S-Curve weekly time-series from an S-Curve sheet.

    Sheet layout (all 5 S-Curve sheets share this structure):
      Row 41  – Ordinal Week dates (datetime), data starts from column 5
      Row 42  – Plan (BL Rev.A1) [A]  Incremental  (skip)
      Row 43  – Plan (BL Rev.A1) [A]  Cumulative   ← use this
      Row 44  – Forecast [C]          Incremental  (skip)
      Row 45  – Forecast [C]          Cumulative   ← use this
      Row 46  – Actual [B]            Incremental  (skip)
      Row 47  – Actual [B]            Cumulative   ← use this

    Rows 12-16 contain the current-week summary snapshot (not a time series).
    """
    if sheet_name not in wb.sheetnames:
        return None

    ws = wb[sheet_name]

    DATE_ROW       = 41
    PLAN_CUM_ROW   = 43
    FORE_CUM_ROW   = 45
    ACT_CUM_ROW    = 47
    PLAN_INCR_ROW  = 42
    FORE_INCR_ROW  = 44
    ACT_INCR_ROW   = 46
    START_COL      = 5  # column E

    # Collect all date columns
    date_cols = []
    for col in range(START_COL, ws.max_column + 1):
        cell_val = ws.cell(row=DATE_ROW, column=col).value
        if cell_val is None:
            continue
        date_str = safe_date(cell_val)
        if date_str:
            date_cols.append((col, date_str))

    # Build time-series data points
    series = []
    for col, date_str in date_cols:
        plan_incr = safe_float(ws.cell(row=PLAN_INCR_ROW, column=col).value)
        plan_cum  = safe_float(ws.cell(row=PLAN_CUM_ROW,  column=col).value)
        fore_incr = safe_float(ws.cell(row=FORE_INCR_ROW, column=col).value)
        fore_cum  = safe_float(ws.cell(row=FORE_CUM_ROW,  column=col).value)
        act_incr  = safe_float(ws.cell(row=ACT_INCR_ROW,  column=col).value)
        act_cum   = safe_float(ws.cell(row=ACT_CUM_ROW,   column=col).value)

        # Skip fully-empty columns
        if all(v is None for v in [plan_cum, fore_cum, act_cum]):
            continue

        series.append({
            "date":          date_str,
            "plan_incr":     plan_incr,
            "plan_cum":      plan_cum,
            "forecast_incr": fore_incr,
            "forecast_cum":  fore_cum,
            "actual_incr":   act_incr,
            "actual_cum":    act_cum,
        })

    # Current-week summary snapshot (scans from row 12 dynamically)
    summary = []
    for row in range(12, 30):
        phase = safe_str(ws.cell(row=row, column=2).value)
        if not phase or phase.startswith("="):
            continue
        plan_incr = safe_float(ws.cell(row=row, column=5).value)
        plan_cum  = safe_float(ws.cell(row=row, column=6).value)
        fore_incr = safe_float(ws.cell(row=row, column=7).value)
        fore_cum  = safe_float(ws.cell(row=row, column=8).value)
        act_incr  = safe_float(ws.cell(row=row, column=9).value)
        act_cum   = safe_float(ws.cell(row=row, column=10).value)

        # Col 11 (K) is [b-a] = Actual Cum - Plan Cum (Non-Rebased Line)
        dev_k = safe_float(ws.cell(row=row, column=11).value)
        dev_non_rebased = dev_k if dev_k is not None else ((act_cum - plan_cum) if (act_cum is not None and plan_cum is not None) else None)
        dev_rebased = (act_cum - fore_cum) if (act_cum is not None and fore_cum is not None) else None

        summary.append({
            "phase":                 phase,
            "plan_incr":             plan_incr,
            "plan_cum":              plan_cum,
            "forecast_incr":         fore_incr,
            "forecast_cum":          fore_cum,
            "actual_incr":           act_incr,
            "actual_cum":            act_cum,
            "dev_non_rebased":       dev_non_rebased,
            "dev_rebased":           dev_rebased,
            "deviation":             dev_non_rebased,
        })
        if "TOTAL" in phase.upper():
            break

    return {
        "label":   label,
        "series":  series,    # weekly time-series for S-Curve chart
        "summary": summary,   # snapshot table (current week)
    }



def extract_scurve_weekly(wb, sheet_name, label):
    """Legacy wrapper – delegates to extract_scurve."""
    return extract_scurve(wb, sheet_name, label)



def compute_delay_and_lookahead(all_documents, cutoff_date_str):
    """
    Compute delayed and 2-week look-ahead documents per strict user requirements:
    1. All tracking date shall refer from Today and Forecast date (fallback to Plan).
    2. Empty tracking date means "not yet issue" / not submitted.
    3. Delay has 2 types:
       3.1 "Delayed": submit date is after Forecast Date.
       3.2 "Delayed": document not yet submitted and today is overdue the Forecast date.
    4. Recheck and count document Delayed for both types.
    5. For 2 weeks lookahead, forecast how many documents will be Delayed (slipping from baseline / high risk).
    """
    if not cutoff_date_str:
        cutoff_date_str = datetime.date.today().isoformat()
        today = datetime.date.today()
    else:
        try:
            today = datetime.date.fromisoformat(cutoff_date_str)
        except (ValueError, TypeError):
            today = datetime.date.today()

    lookahead_end = today + datetime.timedelta(days=14)
    delayed = []
    lookahead = []

    delayed_type1_count = 0
    delayed_type2_count = 0

    for doc in all_documents:
        # Check milestones: for MR/TBE vs standard engineering deliverables
        if doc.get("is_mr_tbe"):
            milestones = [
                {
                    "milestone": "IFA (A1)",
                    "plan": doc.get("ifa_plan"),
                    "forecast": doc.get("ifa_forecast"),
                    "submit_date": doc.get("ifa_submit_date"),
                },
                {
                    "milestone": "IFA (C1)",
                    "plan": doc.get("afc_plan"),
                    "forecast": doc.get("afc_forecast"),
                    "submit_date": doc.get("afc_submit_date"),
                },
            ]
        else:
            milestones = [
                {
                    "milestone": "IFR",
                    "plan": doc.get("ifr_plan"),
                    "forecast": doc.get("ifr_forecast"),
                    "submit_date": doc.get("ifr_submit_date"),
                },
                {
                    "milestone": "IFA",
                    "plan": doc.get("ifa_plan"),
                    "forecast": doc.get("ifa_forecast"),
                    "submit_date": doc.get("ifa_submit_date"),
                },
                {
                    "milestone": "AFC",
                    "plan": doc.get("afc_plan"),
                    "forecast": doc.get("afc_forecast"),
                    "submit_date": doc.get("afc_submit_date"),
                },
            ]

        for ms in milestones:
            plan_str = ms["plan"]
            forecast_str = ms["forecast"]
            submit_str = ms["submit_date"]

            # Skip if both plan and forecast are unavailable
            if (not plan_str or plan_str in ("N/A", "-")) and (not forecast_str or forecast_str in ("N/A", "-")):
                continue

            plan_dt = None
            if plan_str and plan_str not in ("N/A", "-"):
                try:
                    plan_dt = datetime.date.fromisoformat(plan_str)
                except (ValueError, TypeError):
                    pass

            forecast_dt = None
            if forecast_str and forecast_str not in ("N/A", "-"):
                try:
                    forecast_dt = datetime.date.fromisoformat(forecast_str)
                except (ValueError, TypeError):
                    pass

            submit_dt = None
            if submit_str and submit_str not in ("", "N/A", "-"):
                try:
                    submit_dt = datetime.date.fromisoformat(submit_str)
                except (ValueError, TypeError):
                    pass

            # Primary reference date is FORECAST, falling back to PLAN if FORECAST is unavailable
            ref_dt = forecast_dt if forecast_dt is not None else plan_dt
            if ref_dt is None:
                continue

            # Check for slippage: forecast > plan
            is_slipping = False
            if forecast_dt and plan_dt and forecast_dt > plan_dt:
                is_slipping = True

            entry = {
                "doc_no": doc.get("doc_no", ""),
                "title": doc.get("title", ""),
                "discipline": doc.get("discipline", ""),
                "wp": doc.get("wp", ""),
                "milestone": ms["milestone"],
                "plan_date": plan_str,
                "forecast_date": forecast_str,
                "submit_date": submit_str,
                "status": doc.get("status", ""),
                "is_mr_tbe": doc.get("is_mr_tbe", False),
                "is_slipping": is_slipping,
            }

            # Check Type 3.1 vs Type 3.2 Delay
            if submit_dt is not None:
                # Type 3.1: submit date is after Forecast Date
                if submit_dt > ref_dt:
                    entry["delay_days"] = max((submit_dt - ref_dt).days, 1)
                    entry["delay_type"] = "Type 3.1 (Submitted Late)"
                    entry["delay_type_code"] = "3.1"
                    delayed.append(entry)
                    delayed_type1_count += 1
            else:
                # Type 3.2: document not yet submitted and today is overdue the Forecast date
                if ref_dt <= today:
                    entry["delay_days"] = max((today - ref_dt).days, 1)
                    entry["delay_type"] = "Type 3.2 (Not Submitted & Overdue)"
                    entry["delay_type_code"] = "3.2"
                    delayed.append(entry)
                    delayed_type2_count += 1
                elif today < ref_dt <= lookahead_end:
                    # In 2-week look-ahead
                    days_remaining = (ref_dt - today).days
                    entry["days_remaining"] = days_remaining
                    entry["urgency"] = "this_week" if days_remaining <= 7 else "next_week"
                    lookahead.append(entry)

    # Sort: delayed by delay_days desc, lookahead by days_remaining asc
    delayed.sort(key=lambda x: x.get("delay_days", 0), reverse=True)
    lookahead.sort(key=lambda x: x.get("days_remaining", 0))

    # Forecast how many documents in 2-week lookahead will be delayed (slipping from baseline plan)
    lookahead_forecast_delay = sum(1 for item in lookahead if item.get("is_slipping", False))

    return {
        "delayed": delayed,
        "delayed_count": len(delayed),
        "delayed_type1_count": delayed_type1_count,
        "delayed_type2_count": delayed_type2_count,
        "lookahead": lookahead,
        "lookahead_count": len(lookahead),
        "lookahead_forecast_delay": lookahead_forecast_delay,
        "cutoff_date": cutoff_date_str,
        "reference_date": today.isoformat(),
        "lookahead_end": lookahead_end.isoformat(),
    }


# ─── Main Extraction ──────────────────────────────────────────────────────────

def extract_all_data(excel_path=None, force_reload=False):
    """Read the Excel file and extract all dashboard data with memory optimization & disk caching."""
    excel_file = Path(excel_path) if excel_path else find_latest_excel()
    if not excel_file or not excel_file.exists():
        return {"error": "No Excel file found in the project directory."}

    # 1. Try disk cache first unless forced
    if not force_reload:
        cached_data = load_disk_cache(excel_file)
        if cached_data is not None:
            _data_cache["file"] = excel_file.name
            return cached_data

    print(f"  [Data] Extracting from: {excel_file.name}")
    wb = load_trimmed_workbook(excel_file)

    # Summary
    summary = extract_summary(wb, excel_file.name)

    # Weekly Summary
    weekly = extract_weekly_summary(wb, excel_file.name)

    # Documents from all WP sheets
    all_documents = []
    doc_stats = {}
    for wp in WP_SHEETS:
        docs = extract_documents(wb, wp["sheet"], wp["label"])
        all_documents.extend(docs)

    # Documents from Procurement Engineering MR/TBE sheet
    mr_tbe_docs = extract_mr_tbe_documents(wb)
    all_documents.extend(mr_tbe_docs)

    # Stats per WP (combining both standard engineering deliverables and MR/TBE)
    for wp in WP_SHEETS:
        wp_label = wp["label"]
        wp_all = [d for d in all_documents if d["wp"] == wp_label]
        total = len(wp_all)
        submitted = sum(1 for d in wp_all if d["status"] != "Not Yet Submitted")
        not_submitted = total - submitted
        pending_final_approval = sum(1 for d in wp_all if d.get("is_pending_final_approval"))
        doc_stats[wp_label] = {
            "total": total,
            "submitted": submitted,
            "not_submitted": not_submitted,
            "pending_final_approval": pending_final_approval,
            "cy_ap_pending": pending_final_approval,
        }

    pending_final_approval_total = sum(1 for d in all_documents if d.get("is_pending_final_approval"))
    summary["pending_final_approval_count"] = pending_final_approval_total
    summary["cy_ap_pending_count"] = pending_final_approval_total

    # Overdue
    overdue_summary = extract_overdue_summary(wb)
    overdue_list = extract_overdue_list(wb, excel_file.name)

    # S-Curve data
    scurves = {}
    for sc in SCURVE_SHEETS:
        data = extract_scurve(wb, sc["sheet"], sc["label"])
        if data:
            scurves[sc["label"]] = data

    # Delay and look-ahead analysis
    delay_lookahead = compute_delay_and_lookahead(all_documents, summary.get("cutoff_date"))

    # Discipline list for filters
    disciplines = sorted(set(d["discipline"] for d in all_documents if d["discipline"]))

    wb.close()
    del wb
    gc.collect()

    _data_cache["file"] = excel_file.name

    result = {
        "file_name": excel_file.name,
        "summary": summary,
        "weekly_summary": weekly,
        "documents": all_documents,
        "doc_stats": doc_stats,
        "overdue_summary": overdue_summary,
        "overdue_list": overdue_list,
        "scurves": scurves,
        "delay_lookahead": delay_lookahead,
        "disciplines": disciplines,
        "wp_list": [wp["label"] for wp in WP_SHEETS],
        "total_documents": len(all_documents),
        "pending_final_approval_count": pending_final_approval_total,
        "cy_ap_pending_count": pending_final_approval_total,
        "last_updated": datetime.datetime.now().isoformat(),
    }

    # 3. Save to disk cache for instantaneous cold-boot
    save_disk_cache(excel_file, result)

    return result

def generate_excel_bytes(data_list, columns, sheet_title="Export"):
    """
    Generate a styled, professional Excel file in memory and return binary bytes.
    columns is a list of tuples: (field_key, display_label, min_width)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title

    # Header styling: Navy fill, bold white text, centered, thin border
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_thin = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    ws.row_dimensions[1].height = 28

    # Write headers
    for col_idx, (key, label, min_w) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border_thin

    # Data styling
    data_font = Font(name="Calibri", size=10)
    data_align_left = Alignment(horizontal="left", vertical="center")
    data_align_center = Alignment(horizontal="center", vertical="center")

    col_widths = {col_idx: min_w for col_idx, (key, label, min_w) in enumerate(columns, 1)}

    # Write data rows
    for row_idx, row_data in enumerate(data_list, 2):
        ws.row_dimensions[row_idx].height = 20
        for col_idx, (key, label, min_w) in enumerate(columns, 1):
            val = row_data.get(key, "")
            if val is None or val == "":
                val_display = "—"
            elif isinstance(val, float) and "pct" in key:
                val_display = f"{val * 100:.2f}%"
            else:
                val_display = val

            cell = ws.cell(row=row_idx, column=col_idx, value=val_display)
            cell.font = data_font
            cell.border = border_thin

            # Center alignment for dates, codes, numbers, and categories
            if any(term in key for term in ["date", "days", "milestone", "type", "discipline", "urgency", "pct", "variance", "status", "no"]):
                cell.alignment = data_align_center
            else:
                cell.alignment = data_align_left

            # Calculate width based on content length
            val_str = str(val_display)
            if len(val_str) > col_widths[col_idx]:
                col_widths[col_idx] = min(len(val_str) + 3, 50)

    # Set column widths
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(width, 12)

    # Freeze header row
    ws.freeze_panes = "A2"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
    

# ─── HTTP Server ──────────────────────────────────────────────────────────────

class DashboardHandler(http.server.SimpleHTTPRequestHandler):
    """Custom handler for the engineering dashboard."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(BASE_DIR), **kwargs)

    def end_headers(self):
        self.send_header("Cache-Control", "no-cache, no-store, must-revalidate")
        self.send_header("Pragma", "no-cache")
        self.send_header("Expires", "0")
        super().end_headers()

    def do_GET(self):
        path_clean = self.path.split("?")[0].rstrip("/")
        if path_clean == "/api/data":
            self.send_api_data()
        elif path_clean == "/api/refresh":
            self.send_api_refresh()
        elif path_clean in ("/api/export/delayed", "/api/export/delayed.xlsx"):
            self.send_export_delayed()
        elif path_clean in ("/api/export/lookahead", "/api/export/lookahead.xlsx"):
            self.send_export_lookahead()
        elif path_clean in ("/api/export/documents", "/api/export/documents.xlsx"):
            self.send_export_documents()
        elif path_clean in ("", "/", "/index.html"):
            self.path = "/index.html"
            super().do_GET()
        else:
            super().do_GET()

    def do_POST(self):
        if self.path == "/api/upload":
            self.handle_upload()
        else:
            self.send_response(404)
            self.end_headers()

    def send_api_data(self):
        """Return cached data or wait/return loading state."""
        try:
            with _data_cache["lock"]:
                data = _data_cache["data"]
                is_extracting = _data_cache.get("is_extracting", False)

            if data is None:
                if is_extracting:
                    response = {
                        "status": "loading",
                        "message": "Data extraction from Excel is currently in progress. Please wait a moment..."
                    }
                    json_bytes = json.dumps(response).encode("utf-8")
                    self.send_response(202)
                    self.send_header("Content-Type", "application/json; charset=utf-8")
                    self.send_header("Content-Length", len(json_bytes))
                    self.send_header("Access-Control-Allow-Origin", "*")
                    self.end_headers()
                    self.wfile.write(json_bytes)
                    return
                else:
                    with _data_cache["lock"]:
                        _data_cache["is_extracting"] = True
                    data = extract_all_data()
                    with _data_cache["lock"]:
                        _data_cache["data"] = data
                        _data_cache["is_extracting"] = False

            json_bytes = json.dumps(data, ensure_ascii=False, default=str).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", len(json_bytes))
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(json_bytes)
        except Exception as e:
            with _data_cache["lock"]:
                _data_cache["is_extracting"] = False
            import traceback
            traceback.print_exc()
            self._send_error(500, str(e))

    def send_api_refresh(self):
        """Force re-read Excel file."""
        try:
            print("\n  [API] Refreshing data from Excel file...")
            with _data_cache["lock"]:
                _data_cache["is_extracting"] = True
            new_data = extract_all_data(force_reload=True)
            with _data_cache["lock"]:
                _data_cache["data"] = new_data
                _data_cache["is_extracting"] = False
            print(f"  [API] Refresh complete. File: {new_data.get('file_name', '?')}, Docs: {new_data.get('total_documents', '?')}")

            json_bytes = json.dumps(new_data, ensure_ascii=False, default=str).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", len(json_bytes))
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(json_bytes)
        except Exception as e:
            with _data_cache["lock"]:
                _data_cache["is_extracting"] = False
            import traceback
            traceback.print_exc()
            self._send_error(500, str(e))

    def send_export_delayed(self):
        """Export delayed documents to Excel."""
        try:
            with _data_cache["lock"]:
                data = _data_cache["data"]
            if data is None:
                data = extract_all_data()

            delayed = data.get("delay_lookahead", {}).get("delayed", [])
            cutoff = data.get("delay_lookahead", {}).get("cutoff_date", "unknown")
            columns = [
                ("delay_type", "Delay Type", 24),
                ("wp", "Work Package", 16),
                ("discipline", "Discipline", 12),
                ("doc_no", "Document No.", 28),
                ("title", "Document Title", 40),
                ("milestone", "Milestone", 12),
                ("plan_date", "Plan Date", 14),
                ("forecast_date", "Forecast Date", 14),
                ("submit_date", "Submit Date", 14),
                ("delay_days", "Days Overdue", 14),
            ]
            
            excel_bytes = generate_excel_bytes(delayed, columns, sheet_title="Delayed Deliverables")
            filename = f"Delayed_Documents_{cutoff}.xlsx"

            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", f'attachment; filename="{filename}"; filename*=UTF-8\'\'{filename}')
            self.send_header("Content-Length", str(len(excel_bytes)))
            self.end_headers()
            self.wfile.write(excel_bytes)
        except Exception as e:
            import traceback
            traceback.print_exc()
            self._send_error(500, str(e))

    def send_export_documents(self):
        """Export all documents to Excel."""
        try:
            with _data_cache["lock"]:
                data = _data_cache["data"]
            if data is None:
                data = extract_all_data()

            docs = data.get("documents", [])
            cutoff = data.get("delay_lookahead", {}).get("cutoff_date", "unknown")
            columns = [
                ("doc_no", "Document No.", 28),
                ("title", "Document Title", 40),
                ("wp", "Work Package", 16),
                ("discipline", "Discipline", 12),
                ("type", "Type", 12),
                ("status", "Status", 18),
                ("plan_pct", "Plan %", 12),
                ("forecast_pct", "Forecast %", 12),
                ("actual_pct", "Actual %", 12),
                ("variance", "Variance", 12),
                ("ifr_forecast", "IFR Forecast", 14),
                ("ifr_submit_date", "IFR Submit Date", 14),
                ("ifa_forecast", "IFA Forecast", 14),
                ("ifa_submit_date", "IFA Submit Date", 14),
                ("afc_forecast", "AFC Forecast", 14),
                ("afc_submit_date", "AFC Submit Date", 14),
                ("ap_return_code", "Return Code", 14),
                ("ap_return_date", "Return Date", 14),
                ("is_pending_final_approval", "Pending Final Approval", 22),
            ]
            
            excel_bytes = generate_excel_bytes(docs, columns, sheet_title="Master Document Register")
            filename = f"Engineering_Documents_{cutoff}.xlsx"

            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", f'attachment; filename="{filename}"; filename*=UTF-8\'\'{filename}')
            self.send_header("Content-Length", str(len(excel_bytes)))
            self.end_headers()
            self.wfile.write(excel_bytes)
        except Exception as e:
            import traceback
            traceback.print_exc()
            self._send_error(500, str(e))

    def send_export_lookahead(self):
        """Export lookahead documents to Excel."""
        try:
            with _data_cache["lock"]:
                data = _data_cache["data"]
            if data is None:
                data = extract_all_data()

            lookahead_raw = data.get("delay_lookahead", {}).get("lookahead", [])
            cutoff = data.get("delay_lookahead", {}).get("cutoff_date", "unknown")

            # Enrich display fields for urgency and slippage
            lookahead = []
            for item in lookahead_raw:
                enriched = dict(item)
                urgency = item.get("urgency", "")
                enriched["urgency_text"] = "Due This Week" if urgency == "this_week" else "Due Next Week"
                enriched["slippage_text"] = "⚠️ Slipping" if item.get("is_slipping") else "On Track"
                lookahead.append(enriched)

            columns = [
                ("urgency_text", "Urgency", 18),
                ("wp", "Work Package", 16),
                ("discipline", "Discipline", 12),
                ("doc_no", "Document No.", 28),
                ("title", "Document Title", 40),
                ("milestone", "Milestone", 12),
                ("plan_date", "Plan Date", 14),
                ("forecast_date", "Forecast Date", 14),
                ("days_remaining", "Days Left", 14),
                ("slippage_text", "Slippage", 16),
            ]
            
            excel_bytes = generate_excel_bytes(lookahead, columns, sheet_title="2-Week Lookahead")
            filename = f"Lookahead_Documents_{cutoff}.xlsx"

            self.send_response(200)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", f'attachment; filename="{filename}"; filename*=UTF-8\'\'{filename}')
            self.send_header("Content-Length", str(len(excel_bytes)))
            self.end_headers()
            self.wfile.write(excel_bytes)
        except Exception as e:
            import traceback
            traceback.print_exc()
            self._send_error(500, str(e))

    def handle_upload(self):
        """Handle Excel file upload. Saves to BASE_DIR and refreshes data."""
        try:
            content_type = self.headers.get("Content-Type", "")
            if "multipart/form-data" not in content_type:
                self._send_error(400, "Expected multipart/form-data")
                return

            # Parse multipart form data
            boundary = content_type.split("boundary=")[1].strip()
            content_length = int(self.headers.get("Content-Length", 0))
            body = self.rfile.read(content_length)

            # Find the file in the multipart data
            boundary_bytes = ("--" + boundary).encode()
            parts = body.split(boundary_bytes)

            filename = None
            file_data = None

            for part in parts:
                if b"filename=" in part:
                    # Extract filename
                    header_end = part.find(b"\r\n\r\n")
                    if header_end < 0:
                        continue
                    header = part[:header_end].decode("utf-8", errors="replace")
                    fn_match = re.search(r'filename="([^"]+)"', header)
                    if fn_match:
                        filename = fn_match.group(1)
                    file_data = part[header_end + 4:]
                    # Remove trailing boundary markers
                    if file_data.endswith(b"\r\n"):
                        file_data = file_data[:-2]
                    if file_data.endswith(b"--"):
                        file_data = file_data[:-2]
                    if file_data.endswith(b"\r\n"):
                        file_data = file_data[:-2]

            if not filename or not file_data:
                self._send_error(400, "No file found in upload")
                return

            if not filename.endswith(".xlsx"):
                self._send_error(400, "Only .xlsx files are supported")
                return
                
            if "MDR" not in filename.upper():
                self._send_error(400, "Invalid file format. Only MDR Excel files are supported.")
                return

            # Save the uploaded file
            dest = BASE_DIR / filename
            with open(dest, "wb") as f:
                f.write(file_data)

            print(f"\n  [Upload] Saved: {filename} ({len(file_data)} bytes)")

            # Refresh data with the new file
            with _data_cache["lock"]:
                _data_cache["is_extracting"] = True
            new_data = extract_all_data(excel_path=dest, force_reload=True)
            with _data_cache["lock"]:
                _data_cache["data"] = new_data
                _data_cache["is_extracting"] = False
            print(f"  [Upload] Data refreshed. Docs: {new_data.get('total_documents', '?')}")

            response = {
                "success": True,
                "filename": filename,
                "message": f"File uploaded and data refreshed. {new_data.get('total_documents', 0)} documents loaded.",
                "data": new_data,
            }
            json_bytes = json.dumps(response, ensure_ascii=False, default=str).encode("utf-8")
            self.send_response(200)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", len(json_bytes))
            self.send_header("Access-Control-Allow-Origin", "*")
            self.end_headers()
            self.wfile.write(json_bytes)

        except Exception as e:
            with _data_cache["lock"]:
                _data_cache["is_extracting"] = False
            import traceback
            traceback.print_exc()
            self._send_error(500, str(e))

    def _send_error(self, code, message):
        error_msg = json.dumps({"error": message}).encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", len(error_msg))
        self.end_headers()
        self.wfile.write(error_msg)

    def log_message(self, format, *args):
        try:
            msg = str(args[0]) if args else ""
            if "/api/" in msg:
                super().log_message(format, *args)
        except Exception:
            pass


class ThreadedHTTPServer(socketserver.ThreadingMixIn, http.server.HTTPServer):
    """Handle requests in separate threads."""
    daemon_threads = True
    allow_reuse_address = True


# ─── Main ────────────────────────────────────────────────────────────────────

def preload_data():
    print("\n  [Background] Pre-loading Excel data...")
    try:
        with _data_cache["lock"]:
            if _data_cache["data"] is not None:
                return
            _data_cache["is_extracting"] = True

        data = extract_all_data()

        with _data_cache["lock"]:
            _data_cache["data"] = data
            _data_cache["is_extracting"] = False

        total = data.get("total_documents", "?")
        fname = data.get("file_name", "?")
        print(f"  [Background] Loaded: {fname} — {total} documents\n")
    except Exception as e:
        with _data_cache["lock"]:
            _data_cache["is_extracting"] = False
        print(f"  [Background] Warning: Could not pre-load data: {e}\n")


if __name__ == "__main__":
    print("=" * 58)
    print("   Engineering Progress Tracking Dashboard Server")
    print("=" * 58)
    print(f"   URL: http://localhost:{PORT}")
    print("   Press Ctrl+C to stop")
    print("=" * 58)

    excel = find_latest_excel()
    print(f"\n  Excel file: {excel.name if excel else 'NOT FOUND'}")
    if excel:
        dt = parse_date_from_filename(excel.name)
        print(f"  Cut-off date from filename: {dt}")
    print(f"\n  Drop new Excel files in: {BASE_DIR}")
    print("  Or use the Upload button on the dashboard.")

    # Start pre-loading in background thread
    threading.Thread(target=preload_data, daemon=True).start()

    print(f"\n  Server ready at http://localhost:{PORT}\n")

    server = ThreadedHTTPServer(("", PORT), DashboardHandler)
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n\nServer stopped.")
        server.server_close()
